import argparse
import csv
from datetime import datetime
import os
from pathlib import Path
import sys
from typing import Any, Dict, List, Optional, Set, Union

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def _check_docx():
    if not HAS_DOCX:
        raise ImportError(
            "The 'python-docx' package is required for DocForge.\n"
            "Install it via: pip install python-docx"
        )


def _replace_in_paragraph(paragraph: Any, mapping: Dict[str, str]) -> None:
    """
    Replace placeholders in a paragraph while preserving run formatting whenever possible.
    """
    full_text = paragraph.text
    if not any(ph in full_text for ph in mapping):
        return

    for placeholder, replacement in mapping.items():
        if placeholder not in full_text:
            continue

        # Check if the placeholder fits inside an individual run (preserves bold, font, color)
        replaced_in_run = False
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, replacement)
                replaced_in_run = True

        # If placeholder is split across multiple XML runs, fallback to paragraph level
        if not replaced_in_run and placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(placeholder, replacement)


def _replace_in_table(table: Any, mapping: Dict[str, str]) -> None:
    """Recursively replace placeholders across all cells and nested tables in a Word table."""
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                _replace_in_paragraph(para, mapping)
            for nested_table in cell.tables:
                _replace_in_table(nested_table, mapping)


def replace_document_placeholders(doc: Any, mapping: Dict[str, str]) -> None:
    """Replace all dictionary placeholders throughout paragraphs, tables, headers, and footers."""
    for para in doc.paragraphs:
        _replace_in_paragraph(para, mapping)

    for table in doc.tables:
        _replace_in_table(table, mapping)

    for section in doc.sections:
        for header_para in section.header.paragraphs:
            _replace_in_paragraph(header_para, mapping)
        for footer_para in section.footer.paragraphs:
            _replace_in_paragraph(footer_para, mapping)


class DocForge:
    """
    Batch Word (.docx) document generator and template mail-merge engine.
    
    Generates numbered series (e.g. PCB measurement protocols / Messprotokolle)
    or data-driven reports from CSV tables with formatting preservation.
    """

    def __init__(self, template_path: Union[str, Path], output_dir: Union[str, Path]):
        _check_docx()
        self.template_path = Path(template_path)
        self.output_dir = Path(output_dir)

        if not self.template_path.is_file():
            raise FileNotFoundError(f"Template file not found: {self.template_path}")

    def generate_single(
        self,
        mapping: Dict[str, str],
        output_filename: str,
    ) -> Path:
        """Generate a single document by replacing a dictionary of placeholders."""
        self.output_dir.mkdir(parents=True, exist_ok=True)
        doc = Document(str(self.template_path))
        replace_document_placeholders(doc, mapping)

        if not output_filename.endswith(".docx"):
            output_filename = f"{output_filename}.docx"

        out_path = self.output_dir / output_filename
        doc.save(str(out_path))
        return out_path

    def generate_range(
        self,
        start: int,
        end: int,
        exclude: Optional[Set[int]] = None,
        date_code: Optional[Union[str, bool]] = True,
        prefix: str = "Report_",
        placeholder: str = "{SERIAL_NUMBER}",
        extra_mapping: Optional[Dict[str, str]] = None,
        verbose: bool = True,
    ) -> List[Path]:
        """
        Generate a batch of documents across a numeric range (e.g. boards 1 to 109).

        Args:
            start: Starting board/part number (inclusive).
            end: Ending board/part number (inclusive).
            exclude: Optional set or list of numbers to skip (e.g. non-existent boards).
            date_code: Suffix date code (e.g. '240924'). If True, uses today's YYMMDD. If False/None, omitted.
            prefix: Filename prefix (e.g. 'Messprotokoll_').
            placeholder: The placeholder token inside the Word template (default: '{SERIAL_NUMBER}').
            extra_mapping: Additional placeholders to replace (e.g. {'{STATUS}': 'PASS'}).
            verbose: Print progress to console.
        """
        self.output_dir.mkdir(parents=True, exist_ok=True)
        exclude_set = set(exclude) if exclude else set()

        if date_code is True:
            actual_date_code = datetime.now().strftime("%y%m%d")
        elif isinstance(date_code, str):
            actual_date_code = date_code
        else:
            actual_date_code = ""

        generated_files = []

        for num in range(start, end + 1):
            if num in exclude_set:
                continue

            serial_str = f"{num}{actual_date_code}"
            mapping = {placeholder: serial_str}
            if extra_mapping:
                mapping.update(extra_mapping)

            filename = f"{prefix}{serial_str}.docx"
            out_path = self.generate_single(mapping, filename)
            generated_files.append(out_path)

            if verbose:
                print(f"Generated: {out_path.name} (Unit #{num})")

        if verbose:
            print(f"\nDone! Created {len(generated_files)} document(s) in: {self.output_dir.resolve()}")
        return generated_files

    def generate_from_csv(
        self,
        csv_path: Union[str, Path],
        filename_pattern: str = "Report_{SERIAL_NUMBER}.docx",
        verbose: bool = True,
    ) -> List[Path]:
        """
        Generate documents driven by a CSV file. Each row fills one document.
        Column headers match placeholder names (e.g. 'SERIAL_NUMBER' -> '{SERIAL_NUMBER}').
        """
        self.output_dir.mkdir(parents=True, exist_ok=True)
        csv_file = Path(csv_path)
        if not csv_file.is_file():
            raise FileNotFoundError(f"CSV file not found: {csv_file}")

        generated_files = []
        with open(csv_file, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row_idx, row in enumerate(reader, start=1):
                # Build placeholder mapping
                mapping = {}
                for col_name, val in row.items():
                    key = col_name.strip()
                    if not key.startswith("{"):
                        key = f"{{{key}}}"
                    mapping[key] = str(val).strip()

                # Derive filename
                out_name = filename_pattern
                for key, val in mapping.items():
                    out_name = out_name.replace(key, val)

                out_path = self.generate_single(mapping, out_name)
                generated_files.append(out_path)

                if verbose:
                    print(f"Row {row_idx}: {out_path.name}")

        if verbose:
            print(f"\nDone! Generated {len(generated_files)} document(s) from {csv_file.name}")
        return generated_files


def main():
    parser = argparse.ArgumentParser(
        description="DocForge: Batch Word (.docx) report generator with placeholder replacement."
    )
    parser.add_argument("template", type=Path, help="Path to the Word template (.docx) file")
    parser.add_argument("-o", "--out", type=Path, required=True, help="Output directory for generated files")

    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument(
        "--range",
        nargs=2,
        type=int,
        metavar=("START", "END"),
        help="Generate numeric range (e.g. --range 1 109)",
    )
    group.add_argument(
        "--csv",
        type=Path,
        metavar="CSV_FILE",
        help="Generate documents driven by a CSV file",
    )

    parser.add_argument(
        "--exclude",
        nargs="*",
        type=int,
        default=[],
        help="Numbers to exclude in range mode (e.g. --exclude 89 90 107 108 109)",
    )
    parser.add_argument(
        "--prefix",
        default="Messprotokoll_",
        help="Filename prefix for range mode (default: 'Messprotokoll_')",
    )
    parser.add_argument(
        "--date-code",
        default=None,
        help="Date code suffix (e.g. '240924'). If omitted, uses current date.",
    )
    parser.add_argument(
        "--no-date",
        action="store_true",
        help="Do not append a date code to serial numbers",
    )
    parser.add_argument(
        "--placeholder",
        default="{SERIAL_NUMBER}",
        help="Template placeholder to replace (default: '{SERIAL_NUMBER}')",
    )

    args = parser.parse_args()

    _check_docx()
    forge = DocForge(args.template, args.out)

    if args.range:
        start_num, end_num = args.range
        date_setting = False if args.no_date else (args.date_code or True)
        forge.generate_range(
            start=start_num,
            end=end_num,
            exclude=set(args.exclude),
            date_code=date_setting,
            prefix=args.prefix,
            placeholder=args.placeholder,
        )
    elif args.csv:
        forge.generate_from_csv(args.csv)


if __name__ == "__main__":
    main()
